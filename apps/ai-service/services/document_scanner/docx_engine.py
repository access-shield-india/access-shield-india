"""DOCX accessibility scanning engine for the Document Scanner."""

import logging
import re
from typing import Optional

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn

from services.document_scanner.base import (
    BaseDocumentEngine,
    Severity,
    Standard,
    truncate_excerpt,
)
from services.document_scanner.locator import DocxLocator, ParagraphLocation

logger = logging.getLogger(__name__)

GENERIC_LINK_TEXT = frozenset(
    {"click here", "here", "read more", "more", "link", "this link", "click", "www"}
)
MANUAL_BULLET_PREFIXES = ("•", "-", "*", "◦", "▪", "–")
NON_DESCRIPTIVE_ALT_PATTERN = re.compile(
    r"^(image\d*\.(png|jpg|jpeg|gif|bmp|emf|wmf)|picture\s*\d*|image\s*\d*|"
    r"screenshot\s*\d*|logo|graphic\s*\d*|chart\s*\d*|diagram\s*\d*)$",
    re.IGNORECASE,
)
BODY_STYLES = frozenset({"Normal", "Body Text", "Body", "Default Paragraph Font"})

#: Locations listed per aggregated finding. Beyond this the list stops being
#: useful and the reader should fix the pattern rather than chase instances.
MAX_LISTED_OCCURRENCES = 25

#: Indic script ranges, used to recommend a specific BCP 47 language tag
#: rather than telling the author to "set a language". Order matters only for
#: reporting; the first match with the highest character count wins.
SCRIPT_RANGES: tuple[tuple[str, str, str, str], ...] = (
    ("\u0900", "\u097f", "hi-IN", "Hindi (Devanagari)"),
    ("\u0980", "\u09ff", "bn-IN", "Bengali"),
    ("\u0a00", "\u0a7f", "pa-IN", "Punjabi (Gurmukhi)"),
    ("\u0a80", "\u0aff", "gu-IN", "Gujarati"),
    ("\u0b00", "\u0b7f", "or-IN", "Odia"),
    ("\u0b80", "\u0bff", "ta-IN", "Tamil"),
    ("\u0c00", "\u0c7f", "te-IN", "Telugu"),
    ("\u0c80", "\u0cff", "kn-IN", "Kannada"),
    ("\u0d00", "\u0d7f", "ml-IN", "Malayalam"),
)


class DocxAccessibilityEngine(BaseDocumentEngine):
    """Runs WCAG 2.1 AA and GIGW 3.0 checks against a Word document."""

    async def run_all_checks(self) -> list:
        checks = [
            ("document_title", self.check_document_title),
            ("document_language", self.check_document_language),
            ("heading_structure", self.check_heading_structure),
            ("image_alt_text", self.check_image_alt_text),
            ("table_headers", self.check_table_headers),
            ("colour_only_content", self.check_colour_only_content),
            ("link_text", self.check_link_text),
            ("list_structure", self.check_list_structure),
            ("form_fields", self.check_form_fields),
        ]
        try:
            self.doc = Document(self.file_path)
        except Exception as e:
            logger.error("Failed to open DOCX: %s", e, exc_info=True)
            self.add_scan_error("document_load", e)
            return self.violations

        try:
            self.locator = DocxLocator(self.doc)
        except Exception as e:
            logger.warning("DOCX locator unavailable, falling back to paragraph indexes: %s", e)
            self.locator = None

        for name, check_fn in checks:
            try:
                await check_fn()
            except Exception as e:
                logger.error("DOCX check '%s' failed: %s", name, e, exc_info=True)
                self.add_scan_error(name, e)
        return self.violations

    # ── Location helpers ────────────────────────────────────────────────────

    def _locate(self, paragraph_index: Optional[int]) -> Optional[ParagraphLocation]:
        if paragraph_index is None or self.locator is None:
            return None
        try:
            return self.locator.locate(paragraph_index)
        except Exception:
            return None

    def _anchor_fields(self, paragraph_index: Optional[int], excerpt: str = "") -> dict:
        """Anchor keyword arguments for `add_violation`, for one paragraph."""
        location = self._locate(paragraph_index)
        if location is None:
            index_label = (
                f"Paragraph {paragraph_index + 1}"
                if paragraph_index is not None
                else "Document body"
            )
            return {
                "location": index_label,
                "paragraph": (paragraph_index + 1) if paragraph_index is not None else None,
                "excerpt": excerpt,
            }

        return {
            "location": location.describe(),
            "page": location.page,
            "line": location.line,
            "paragraph": location.paragraph,
            "heading_path": location.heading_path or None,
            "excerpt": excerpt,
        }

    def _occurrence(self, paragraph_index: Optional[int], excerpt: str = "") -> dict:
        """One entry for an aggregated finding's `occurrence_list`."""
        location = self._locate(paragraph_index)
        if location is None:
            return {
                "paragraph": (paragraph_index + 1) if paragraph_index is not None else None,
                "excerpt": truncate_excerpt(excerpt),
            }
        return {
            "page": location.page,
            "line": location.line,
            "paragraph": location.paragraph,
            "heading_path": location.heading_path or None,
            "excerpt": truncate_excerpt(excerpt),
        }

    def _aggregate_location(self, occurrences: list[dict]) -> str:
        """Summary location string for an aggregated finding."""
        if not occurrences:
            return "Throughout the document"

        pages = sorted({o["page"] for o in occurrences if o.get("page")})
        if pages:
            if len(pages) == 1:
                return f"Page {pages[0]}"
            return f"{len(occurrences)} places, pages {pages[0]}–{pages[-1]}"
        return f"{len(occurrences)} places in the document body"

    # ── Document-level metadata ─────────────────────────────────────────────

    def _first_heading_text(self) -> Optional[str]:
        """
        The document's own headline, used to propose a real title instead of
        telling the author to invent one.
        """
        for paragraph in self.doc.paragraphs:
            style_name = (paragraph.style.name if paragraph.style else "") or ""
            text = (paragraph.text or "").strip()
            if not text:
                continue
            if style_name.lower() == "title" or style_name.startswith("Heading"):
                return text
        # No styled heading: fall back to the first substantial line of text.
        for paragraph in self.doc.paragraphs:
            text = (paragraph.text or "").strip()
            if len(text) >= 10:
                return text
        return None

    def _document_text_sample(self, limit: int = 20000) -> str:
        parts: list[str] = []
        total = 0
        for paragraph in self.doc.paragraphs:
            text = paragraph.text or ""
            if not text:
                continue
            parts.append(text)
            total += len(text)
            if total >= limit:
                break
        return "\n".join(parts)

    def _detect_language(self) -> tuple[str, str]:
        """
        Recommend a BCP 47 tag for this document based on the scripts actually
        present, so the fix names a specific value.
        """
        sample = self._document_text_sample()
        best_tag, best_name, best_count = "en-IN", "Indian English", 0

        for start, end, tag, name in SCRIPT_RANGES:
            count = sum(1 for ch in sample if start <= ch <= end)
            if count > best_count:
                best_tag, best_name, best_count = tag, name, count

        # Require a meaningful amount of the script before recommending it, so a
        # stray glyph does not flip the whole document's language.
        if best_count < 20:
            return "en-IN", "Indian English"
        return best_tag, best_name

    async def check_document_title(self) -> None:
        title = self.doc.core_properties.title
        if title and str(title).strip():
            return

        suggested = self._first_heading_text()
        if suggested:
            suggestion_step = (
                f'Type the document title. Based on this file\'s own headline, '
                f'"{suggested[:120]}" is likely correct.'
            )
            description = (
                "The document properties contain no Title, even though the document "
                f'opens with "{suggested[:80]}".'
            )
        else:
            suggestion_step = (
                "Type the full official name of the document, as it would appear on a "
                "covering letter."
            )
            description = "The document properties contain no Title."

        self.add_violation(
            violation_id="docx_no_title_001",
            checkpoint_id="GIGW_5.2.28",
            standard=Standard.WCAG_2_1_AA,
            severity=Severity.SERIOUS,
            category="metadata",
            wcag_criterion="2.4.2",
            description=description,
            location="File → Info → Properties → Title",
            excerpt=suggested or "",
            impact=(
                "Screen readers announce the file name or 'Untitled Document' when the "
                "document opens, so a reader cannot tell what they have opened without "
                "reading into the body. The empty Title also carries over into the PDF "
                "if this file is exported."
            ),
            remediation=(
                "Set the Title in document properties, then re-save. "
                "In Word: File → Info → Properties panel on the right → Title."
            ),
            fix_steps=[
                "In Word, open File → Info.",
                "In the Properties panel on the right, click the Title field.",
                suggestion_step,
                "Save the document. If you export to PDF, the Title carries across.",
            ],
            auto_fixable=True,
        )

    async def check_document_language(self) -> None:
        language = self.doc.core_properties.language
        if language and str(language).strip():
            return

        tag, name = self._detect_language()

        self.add_violation(
            violation_id="docx_no_language_001",
            checkpoint_id="GIGW_5.2.38",
            standard=Standard.WCAG_2_1_AA,
            severity=Severity.SERIOUS,
            category="language",
            wcag_criterion="3.1.1",
            description=(
                "The document does not declare a language. Based on the text in this "
                f"file the correct value is {tag} ({name})."
            ),
            location="Review → Language → Set Proofing Language",
            excerpt=tag,
            impact=(
                "A screen reader falls back to the language of the user's operating "
                "system. English text read with a Hindi voice profile, or Hindi text "
                "read with an English one, is close to unintelligible."
            ),
            remediation=(
                f"Set the document language to {name} ({tag}). "
                "In Word: select all text (Ctrl+A), then Review → Language → "
                "Set Proofing Language."
            ),
            fix_steps=[
                "Select the whole document with Ctrl+A (Cmd+A on macOS).",
                "Open Review → Language → Set Proofing Language.",
                f"Choose {name} and confirm the tag reads {tag}.",
                (
                    "Tick 'Set As Default' so new content inherits it, then click OK "
                    "and save."
                ),
                (
                    "If any passage is in a different language, select just that passage "
                    "and repeat these steps for it."
                ),
            ],
        )

    # ── Structure ───────────────────────────────────────────────────────────

    async def check_heading_structure(self) -> None:
        headings: list[tuple[int, int, str]] = []
        word_count = 0

        for para_idx, para in enumerate(self.doc.paragraphs):
            word_count += len(para.text.split())
            style_name = para.style.name if para.style else ""
            if style_name.startswith("Heading"):
                try:
                    level = int(style_name.split(" ")[-1])
                except ValueError:
                    continue
                headings.append((para_idx, level, para.text.strip()))

        if not headings and word_count > 500:
            candidates = self._bold_pseudo_headings()
            occurrences = [
                self._occurrence(idx, text) for idx, text in candidates[:MAX_LISTED_OCCURRENCES]
            ]
            self.add_violation(
                violation_id="docx_no_headings_001",
                checkpoint_id="GIGW_5.2.7",
                standard=Standard.WCAG_2_1_AA,
                severity=Severity.SERIOUS,
                category="heading_structure",
                wcag_criterion="1.3.1",
                description=(
                    f"The document contains about {word_count} words but no Heading "
                    "styles, so it has no navigable structure."
                    + (
                        f" {len(candidates)} paragraph(s) look like headings — short, "
                        "bold lines — but use body text styling."
                        if candidates
                        else ""
                    )
                ),
                location=(
                    "Entire document"
                    if not candidates
                    else self._aggregate_location(occurrences)
                ),
                impact=(
                    "Screen reader users navigate long documents by jumping between "
                    "headings. Without them the only way through this document is to "
                    "read every paragraph in order from the beginning."
                ),
                remediation=(
                    "Apply Word's built-in Heading styles to the section titles. "
                    "Bold text is styling only — assistive technology cannot see it as "
                    "a heading."
                ),
                fix_steps=[
                    (
                        "Click into the first section title, then choose Heading 1 from "
                        "Home → Styles."
                    ),
                    (
                        "Apply Heading 2 to its subsections and Heading 3 below those, "
                        "without skipping a level."
                    ),
                    (
                        "Check the result in View → Navigation Pane. The pane should read "
                        "like a table of contents."
                    ),
                    (
                        "Remove the manual bold-and-enlarge formatting that was standing "
                        "in for headings."
                    ),
                ],
                occurrences=max(len(candidates), 1),
                occurrence_list=occurrences,
            )

        empty_headings: list[dict] = []
        skipped_levels: list[tuple[int, int, int, str]] = []

        prev_level = 0
        for para_idx, level, text in headings:
            if not text:
                empty_headings.append(
                    self._occurrence(para_idx, self._following_text(para_idx))
                )
            elif prev_level > 0 and level > prev_level + 1:
                skipped_levels.append((para_idx, prev_level, level, text))

            if text:
                prev_level = level

        if empty_headings:
            self.add_violation(
                violation_id="docx_empty_heading_001",
                checkpoint_id="GIGW_5.2.7",
                standard=Standard.WCAG_2_1_AA,
                severity=Severity.MODERATE,
                category="heading_structure",
                wcag_criterion="2.4.6",
                description=(
                    f"{len(empty_headings)} paragraph(s) carry a Heading style but "
                    "contain no text."
                ),
                location=self._aggregate_location(empty_headings),
                impact=(
                    "Empty headings appear in a screen reader's heading list as blank "
                    "entries. A reader jumping through the document lands on nothing and "
                    "cannot tell whether content is missing or the file is broken."
                ),
                remediation=(
                    "Delete the empty heading paragraphs, or give each one the section "
                    "title it was meant to introduce."
                ),
                fix_steps=[
                    (
                        "Turn on Home → Show/Hide (¶) so empty paragraphs are visible."
                    ),
                    (
                        "Go to each location listed above. The 'Text to search for' column "
                        "shows the content that follows the empty heading."
                    ),
                    (
                        "If the blank line exists only for spacing, delete it and use "
                        "Paragraph → Spacing Before/After instead."
                    ),
                    (
                        "If it should introduce the section below it, type that section's "
                        "title into it."
                    ),
                ],
                occurrences=len(empty_headings),
                occurrence_list=empty_headings[:MAX_LISTED_OCCURRENCES],
                auto_fixable=True,
            )

        for index, (para_idx, from_level, to_level, text) in enumerate(skipped_levels):
            anchors = self._anchor_fields(para_idx, text)
            self.add_violation(
                violation_id=f"docx_heading_skip_{index}",
                checkpoint_id="GIGW_5.2.7",
                standard=Standard.WCAG_2_1_AA,
                severity=Severity.MODERATE,
                category="heading_structure",
                wcag_criterion="1.3.1",
                description=(
                    f'The heading "{text[:80]}" is styled Heading {to_level}, but the '
                    f"heading before it is Heading {from_level}. Heading "
                    f"{from_level + 1} is missing, so the outline jumps a level."
                ),
                impact=(
                    "A screen reader announces heading levels as it moves. Jumping from "
                    f"level {from_level} to level {to_level} implies a subsection that "
                    "does not exist, so readers cannot tell how this section relates to "
                    "the one above it."
                ),
                remediation=(
                    f'Change "{text[:60]}" from Heading {to_level} to '
                    f"Heading {from_level + 1}."
                ),
                fix_steps=[
                    f'Go to the heading "{text[:80]}".',
                    (
                        f"Apply Heading {from_level + 1} from Home → Styles instead of "
                        f"Heading {to_level}."
                    ),
                    (
                        "If the current level is deliberate, add the missing "
                        f"Heading {from_level + 1} section above it instead."
                    ),
                    "Confirm the outline in View → Navigation Pane.",
                ],
                auto_fixable=True,
                **anchors,
            )

    def _bold_pseudo_headings(self) -> list[tuple[int, str]]:
        """Short, fully-bold body paragraphs — headings faked with formatting."""
        found: list[tuple[int, str]] = []
        for para_idx, para in enumerate(self.doc.paragraphs):
            text = (para.text or "").strip()
            if not text or len(text) > 100:
                continue
            style_name = (para.style.name if para.style else "") or ""
            if style_name.startswith("Heading") or style_name.lower() == "title":
                continue
            runs = [r for r in para.runs if (r.text or "").strip()]
            if runs and all(r.bold for r in runs):
                found.append((para_idx, text))
        return found

    def _following_text(self, paragraph_index: int) -> str:
        """First non-empty paragraph after the given index, for context."""
        paragraphs = self.doc.paragraphs
        for idx in range(paragraph_index + 1, min(paragraph_index + 6, len(paragraphs))):
            text = (paragraphs[idx].text or "").strip()
            if text:
                return text
        return ""

    # ── Images ──────────────────────────────────────────────────────────────

    def _inline_shape_paragraphs(self) -> list[int]:
        """Paragraph index for each inline drawing, in document order."""
        indexes: list[int] = []
        try:
            for para_idx, para in enumerate(self.doc.paragraphs):
                for drawing in para._p.findall(f".//{qn('w:drawing')}"):
                    if drawing.find(f".//{qn('wp:inline')}") is not None:
                        indexes.append(para_idx)
        except Exception as e:
            logger.warning("Could not map inline shapes to paragraphs: %s", e)
        return indexes

    def _get_inline_shape_alt(self, shape) -> tuple[str, str]:
        """Return (descr, name) from inline shape XML."""
        shape_id = getattr(shape, "shape_id", 0)
        shape_name = f"Image {shape_id}"
        descr = ""

        try:
            inline = shape._inline
            for elem in inline.iter():
                if elem.tag.endswith("}cNvPr"):
                    descr = elem.get("descr", "") or ""
                    shape_name = elem.get("name", shape_name) or shape_name
                    break
        except Exception:
            pass
        return descr, shape_name

    async def check_image_alt_text(self) -> None:
        paragraph_indexes = self._inline_shape_paragraphs()

        for position, shape in enumerate(self.doc.inline_shapes):
            descr, shape_name = self._get_inline_shape_alt(shape)
            shape_id = getattr(shape, "shape_id", id(shape))
            para_idx = (
                paragraph_indexes[position] if position < len(paragraph_indexes) else None
            )
            location = self._locate(para_idx)
            nearby = self._nearby_text(para_idx)

            # Where the image sits matters more than its internal name, which
            # authors never see, so lead with position and surrounding text.
            position_label = (
                location.describe() if location else f"Image {position + 1} in document order"
            )
            anchors = self._anchor_fields(para_idx, nearby or shape_name)
            anchors["location"] = f"{position_label} — image \"{shape_name}\""

            if not descr.strip():
                self.add_violation(
                    violation_id=f"docx_no_alt_{shape_id}",
                    checkpoint_id="GIGW_5.2.1",
                    standard=Standard.WCAG_2_1_AA,
                    severity=Severity.CRITICAL,
                    category="alt_text",
                    wcag_criterion="1.1.1",
                    description=(
                        f"Image {position + 1} of {len(self.doc.inline_shapes)} "
                        f'("{shape_name}") has no alternative text.'
                        + (f' It sits next to the text "{nearby[:70]}".' if nearby else "")
                    ),
                    impact=(
                        "A screen reader skips this image entirely or announces only "
                        "\"image\". Whatever it shows — a chart, an org structure, a "
                        "signature block, a screenshot — is unavailable to blind readers."
                    ),
                    remediation=(
                        "Right-click the image → Edit Alt Text, then describe what the "
                        "image conveys in context. If it is purely decorative, tick "
                        "'Mark as decorative' instead."
                    ),
                    fix_steps=[
                        f"Go to {position_label} and select the image.",
                        "Right-click it and choose Edit Alt Text.",
                        (
                            "Write one or two sentences saying what the image tells the "
                            "reader, not what it looks like. For a chart, state the "
                            "finding: \"Revenue grew 42% between Q2 and Q3 2024\"."
                        ),
                        (
                            "If the image is a divider, watermark or background flourish "
                            "that adds nothing, tick 'Mark as decorative' instead."
                        ),
                        (
                            "Do not start the text with \"Image of\" — screen readers "
                            "already announce that it is an image."
                        ),
                    ],
                    auto_fixable=False,
                    **anchors,
                )
            elif NON_DESCRIPTIVE_ALT_PATTERN.match(descr.strip()):
                self.add_violation(
                    violation_id=f"docx_nondescriptive_alt_{shape_id}",
                    checkpoint_id="GIGW_5.2.1",
                    standard=Standard.WCAG_2_1_AA,
                    severity=Severity.MODERATE,
                    category="alt_text",
                    wcag_criterion="1.1.1",
                    description=(
                        f'Image {position + 1} has placeholder alternative text: '
                        f'"{descr.strip()}". This is the file name or a generic word, '
                        "not a description."
                    ),
                    impact=(
                        "A screen reader reads the alt text out verbatim. Hearing "
                        f'"{descr.strip()}" tells a blind reader nothing about the '
                        "content, while the presence of alt text stops automated tools "
                        "flagging it — so the gap goes unnoticed."
                    ),
                    remediation=(
                        "Replace the placeholder with a description of what the image "
                        "conveys in this document."
                    ),
                    fix_steps=[
                        f"Go to {position_label} and select the image.",
                        "Right-click it and choose Edit Alt Text.",
                        f'Delete "{descr.strip()}".',
                        (
                            "Describe what the image tells the reader in this context, "
                            "in one or two sentences."
                        ),
                    ],
                    auto_fixable=False,
                    **anchors,
                )

    def _nearby_text(self, paragraph_index: Optional[int]) -> str:
        """Text around a paragraph, used to describe where an image sits."""
        if paragraph_index is None:
            return ""
        paragraphs = self.doc.paragraphs
        own = (paragraphs[paragraph_index].text or "").strip() if (
            0 <= paragraph_index < len(paragraphs)
        ) else ""
        if own:
            return own
        for offset in (1, -1, 2, -2):
            idx = paragraph_index + offset
            if 0 <= idx < len(paragraphs):
                text = (paragraphs[idx].text or "").strip()
                if text:
                    return text
        return ""

    # ── Tables ──────────────────────────────────────────────────────────────

    def _table_paragraph_indexes(self) -> list[Optional[int]]:
        """Nearest preceding body paragraph index for each table."""
        indexes: list[Optional[int]] = []
        try:
            paragraph_count = 0
            for child in self.doc.element.body.iterchildren():
                if child.tag == qn("w:p"):
                    paragraph_count += 1
                elif child.tag == qn("w:tbl"):
                    indexes.append(paragraph_count - 1 if paragraph_count else None)
        except Exception as e:
            logger.warning("Could not map tables to paragraphs: %s", e)
        return indexes

    async def check_table_headers(self) -> None:
        table_paragraphs = self._table_paragraph_indexes()

        for table_num, table in enumerate(self.doc.tables, 1):
            if len(table.rows) < 2:
                continue

            tr_pr = table.rows[0]._tr.find(qn("w:trPr"))
            has_header = tr_pr is not None and tr_pr.find(qn("w:tblHeader")) is not None
            if has_header:
                continue

            first_row = [
                (cell.text or "").strip() for cell in table.rows[0].cells
            ]
            header_preview = " | ".join(v for v in first_row if v)[:150]
            para_idx = (
                table_paragraphs[table_num - 1]
                if table_num - 1 < len(table_paragraphs)
                else None
            )
            location = self._locate(para_idx)
            position_label = (
                f"{location.describe()} — table {table_num}"
                if location
                else f"Table {table_num} in document order"
            )

            anchors = self._anchor_fields(para_idx, header_preview)
            anchors["location"] = position_label

            self.add_violation(
                violation_id=f"docx_table_no_header_{table_num}",
                checkpoint_id="GIGW_5.1.19",
                standard=Standard.WCAG_2_1_AA,
                severity=Severity.SERIOUS,
                category="table_structure",
                wcag_criterion="1.3.1",
                description=(
                    f"Table {table_num} ({len(table.rows)} rows × "
                    f"{len(table.columns)} columns) has no header row marked for "
                    "accessibility."
                    + (
                        f' Its first row reads: "{header_preview}".'
                        if header_preview
                        else ""
                    )
                ),
                impact=(
                    "A screen reader reads a data table cell by cell. Without a marked "
                    "header row it announces bare values with no column name, so a "
                    "reader hears \"4,50,000\" without hearing which column or period it "
                    "belongs to."
                ),
                remediation=(
                    "Mark the first row as a header row so its labels are announced with "
                    "every data cell. Click inside the first row, then Table Design → "
                    "tick Header Row."
                ),
                fix_steps=[
                    f"Go to {position_label} and click inside its first row.",
                    (
                        "On the Table Design tab, tick 'Header Row' in the Table Style "
                        "Options group."
                    ),
                    (
                        "On the Layout tab, click 'Repeat Header Rows' so the header "
                        "repeats if the table breaks across pages."
                    ),
                    (
                        "Confirm the first row holds column labels rather than data. If "
                        "the table has no labels, add a row and name each column."
                    ),
                    (
                        "If this table is only there to position content side by side "
                        "rather than to present data, replace it with columns or a text "
                        "box — layout tables confuse screen readers."
                    ),
                ],
                **anchors,
            )

    # ── Colour ──────────────────────────────────────────────────────────────

    async def check_colour_only_content(self) -> None:
        """
        Flag text that may rely on colour alone (WCAG 1.4.1).

        Only selective emphasis counts. A paragraph that is entirely one colour
        is styling, not information encoding, and flagging every such run
        produced hundreds of false positives that buried the real findings.
        Findings are aggregated per colour so one decision covers one entry.
        """
        highlight_groups: dict[str, list[dict]] = {}
        colour_groups: dict[str, list[dict]] = {}

        for para_idx, para in enumerate(self.doc.paragraphs):
            style_name = (para.style.name if para.style else "") or ""
            # Branded heading colours are a design choice, not a data encoding.
            if style_name.startswith("Heading") or style_name.lower() == "title":
                continue

            runs = [r for r in para.runs if (r.text or "").strip()]
            if not runs:
                continue

            highlighted = [r for r in runs if self._highlight_name(r)]
            coloured = [r for r in runs if self._font_colour_hex(r)]

            # Highlighting is nearly always semantic, so report it regardless of
            # whether the whole paragraph is highlighted.
            for run in highlighted:
                name = self._highlight_name(run) or "highlight"
                highlight_groups.setdefault(name, []).append(
                    self._occurrence(para_idx, run.text)
                )

            # Font colour only signals meaning when part of the paragraph is
            # coloured and part is not.
            if coloured and len(coloured) < len(runs):
                for run in coloured:
                    hex_value = self._font_colour_hex(run) or "000000"
                    colour_groups.setdefault(hex_value, []).append(
                        self._occurrence(para_idx, run.text)
                    )

        for name, occurrences in sorted(highlight_groups.items()):
            self._add_colour_finding(
                group_id=f"highlight_{name}",
                label=f"{name.replace('_', ' ').lower()} highlighting",
                occurrences=occurrences,
                is_highlight=True,
            )

        for hex_value, occurrences in sorted(colour_groups.items()):
            self._add_colour_finding(
                group_id=f"colour_{hex_value}",
                label=f"text coloured #{hex_value}",
                occurrences=occurrences,
                is_highlight=False,
            )

    @staticmethod
    def _highlight_name(run) -> Optional[str]:
        try:
            highlight = run.font.highlight_color
        except Exception:
            return None
        if highlight is None:
            return None
        if highlight in (WD_COLOR_INDEX.AUTO, WD_COLOR_INDEX.WHITE):
            return None
        name = getattr(highlight, "name", None)
        return str(name) if name else str(highlight)

    @staticmethod
    def _font_colour_hex(run) -> Optional[str]:
        try:
            colour = run.font.color
            if colour is None or colour.rgb is None:
                return None
            rgb = colour.rgb
            if (rgb[0], rgb[1], rgb[2]) == (0, 0, 0):
                return None
            return f"{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"
        except Exception:
            return None

    def _add_colour_finding(
        self, group_id: str, label: str, occurrences: list[dict], is_highlight: bool
    ) -> None:
        samples = [o.get("excerpt") for o in occurrences[:3] if o.get("excerpt")]
        sample_text = "; ".join(f'"{s}"' for s in samples)
        mechanism = "highlighting" if is_highlight else "font colour"

        self.add_violation(
            violation_id=f"docx_colour_only_{group_id}",
            checkpoint_id="GIGW_5.2.12",
            standard=Standard.WCAG_2_1_AA,
            severity=Severity.MODERATE,
            category="colour_contrast",
            wcag_criterion="1.4.1",
            description=(
                f"{len(occurrences)} passage(s) use {label} to stand out from the "
                "surrounding text. If that emphasis carries meaning, the meaning is "
                "available only to readers who can see the colour."
                + (f" For example: {sample_text}." if sample_text else "")
            ),
            location=self._aggregate_location(occurrences),
            impact=(
                "Readers with colour blindness, readers using a screen reader, and "
                "anyone reading a monochrome print-out receive the text without the "
                f"{mechanism}. If the {mechanism} marks something as mandatory, "
                "amended, at risk or superseded, that status is lost."
            ),
            remediation=(
                f"Decide whether the {mechanism} carries meaning. If it does, add a "
                "text or symbol cue alongside it. If it is decorative, leave it and "
                "record the decision."
            ),
            fix_steps=[
                (
                    f"Review the passages listed above and decide what the {mechanism} "
                    "is telling the reader."
                ),
                (
                    "If it is purely decorative — brand styling, visual variety — no "
                    "change is needed. Note the decision so the next audit does not "
                    "re-raise it."
                ),
                (
                    "If it signals status, add the status in words as well, for example "
                    '"Mandatory —", "[Amended]", "(At risk)", or a bold label.'
                ),
                (
                    f"If the {mechanism} replaces a table column such as Status or "
                    "Priority, add that column instead."
                ),
                (
                    "Check the colour still meets 4.5:1 contrast against its background "
                    "using the TPGi Colour Contrast Analyser."
                ),
            ],
            occurrences=len(occurrences),
            occurrence_list=occurrences[:MAX_LISTED_OCCURRENCES],
        )

    # ── Links ───────────────────────────────────────────────────────────────

    def _iter_hyperlinks(self):
        """Yield (paragraph_index, link_text, target) for each hyperlink."""
        try:
            rels = self.doc.part.rels
        except Exception:
            rels = {}

        for para_idx, para in enumerate(self.doc.paragraphs):
            for child in para._element:
                if child.tag != qn("w:hyperlink"):
                    continue
                text = "".join(
                    node.text for node in child.iter(qn("w:t")) if node.text
                )
                if not text.strip():
                    continue

                target = ""
                rel_id = child.get(qn("r:id"))
                if rel_id and rel_id in rels:
                    try:
                        target = rels[rel_id].target_ref or ""
                    except Exception:
                        target = ""
                yield para_idx, text.strip(), target

    async def check_link_text(self) -> None:
        bare_urls: list[dict] = []

        for idx, (para_idx, link_text, target) in enumerate(self._iter_hyperlinks()):
            lowered = link_text.lower()
            is_generic = lowered in GENERIC_LINK_TEXT
            is_bare_url = bool(re.match(r"^(https?://|www\.)", lowered))

            if is_bare_url and not is_generic:
                bare_urls.append(self._occurrence(para_idx, link_text))
                continue
            if not is_generic:
                continue

            anchors = self._anchor_fields(para_idx, link_text)
            destination = target or "the linked page"
            self.add_violation(
                violation_id=f"docx_generic_link_{idx}",
                checkpoint_id="GIGW_5.2.30",
                standard=Standard.WCAG_2_1_AA,
                severity=Severity.SERIOUS,
                category="link_text",
                wcag_criterion="2.4.4",
                description=(
                    f'A hyperlink reads "{link_text}", which does not say where it '
                    "goes."
                    + (f" It points to {target}." if target else "")
                ),
                impact=(
                    "Screen reader users often pull up a list of every link on a page "
                    "and tab through it out of context. A list of entries all reading "
                    f'"{link_text}" gives no way to choose between them.'
                ),
                remediation=(
                    f'Replace "{link_text}" with wording that names the destination, '
                    "so the link makes sense read on its own."
                ),
                fix_steps=[
                    f'Find the link that reads "{link_text}" at the location above.',
                    (
                        f"Work out what it points at — {destination} — and what a reader "
                        "would get from following it."
                    ),
                    (
                        "Select the link text and retype it as the destination's name, "
                        'for example "Download the GIGW 3.0 guidelines (PDF, 2.4 MB)".'
                    ),
                    (
                        "Keep the surrounding sentence readable — the link text should be "
                        "part of the sentence, not appended to it."
                    ),
                ],
                **anchors,
            )

        if bare_urls:
            self.add_violation(
                violation_id="docx_bare_url_links_001",
                checkpoint_id="GIGW_5.2.30",
                standard=Standard.WCAG_2_1_AA,
                severity=Severity.MODERATE,
                category="link_text",
                wcag_criterion="2.4.4",
                description=(
                    f"{len(bare_urls)} hyperlink(s) display the raw web address instead "
                    "of descriptive text."
                ),
                location=self._aggregate_location(bare_urls),
                impact=(
                    "A screen reader reads a raw URL character by character, including "
                    "every slash and hyphen. A long address can take half a minute to "
                    "announce and is almost impossible to follow by ear."
                ),
                remediation=(
                    "Replace each raw address with the name of the page it opens, "
                    "keeping the same link target."
                ),
                fix_steps=[
                    "Go to each location listed above.",
                    (
                        "Select the link, right-click and choose Edit Hyperlink (or press "
                        "Ctrl+K)."
                    ),
                    (
                        "Leave the Address unchanged and replace 'Text to display' with "
                        "the page or document name."
                    ),
                    (
                        "If the document is meant to be printed and readers need the "
                        "address, keep the descriptive link and put the full URL in a "
                        "footnote or an appendix."
                    ),
                ],
                occurrences=len(bare_urls),
                occurrence_list=bare_urls[:MAX_LISTED_OCCURRENCES],
            )

    # ── Lists ───────────────────────────────────────────────────────────────

    async def check_list_structure(self) -> None:
        manual_lists: list[dict] = []

        for para_idx, para in enumerate(self.doc.paragraphs):
            style_name = (para.style.name if para.style else "") or ""
            if style_name not in BODY_STYLES:
                continue

            text = (para.text or "").lstrip()
            if not text:
                continue

            starts_with_bullet = text[0] in MANUAL_BULLET_PREFIXES
            starts_with_number = (
                len(text) > 2 and text[0].isdigit() and text[1] in ".)"
            )
            if starts_with_bullet or starts_with_number:
                manual_lists.append(self._occurrence(para_idx, text))

        if not manual_lists:
            return

        self.add_violation(
            violation_id="docx_manual_list_001",
            checkpoint_id="GIGW_5.2.7",
            standard=Standard.WCAG_2_1_AA,
            severity=Severity.MODERATE,
            category="heading_structure",
            wcag_criterion="1.3.1",
            description=(
                f"{len(manual_lists)} paragraph(s) start with a bullet character or a "
                "number typed by hand rather than using Word's list formatting."
            ),
            location=self._aggregate_location(manual_lists),
            impact=(
                "Word's real lists make a screen reader announce \"list of 7 items, "
                "item 1 of 7\", so a reader knows how much there is and where they are "
                "in it. A typed bullet is just a character: the reader hears the symbol "
                "read out, or nothing at all, with no sense of the list's size."
            ),
            remediation=(
                "Convert these paragraphs to real lists with Home → Bullets or "
                "Home → Numbering, and delete the typed characters."
            ),
            fix_steps=[
                "Select the block of paragraphs at each location listed above.",
                (
                    "Click Home → Bullets for unordered lists, or Home → Numbering for "
                    "ordered ones."
                ),
                (
                    "Delete the bullet characters and manual numbers you typed — Word "
                    "now supplies them."
                ),
                (
                    "Use Increase Indent for sub-items so the nesting is recorded in the "
                    "file rather than faked with tabs or spaces."
                ),
                (
                    "Where the numbers matter, such as clause references, check they "
                    "still read correctly after Word takes over numbering."
                ),
            ],
            occurrences=len(manual_lists),
            occurrence_list=manual_lists[:MAX_LISTED_OCCURRENCES],
        )

    # ── Forms ───────────────────────────────────────────────────────────────

    async def check_form_fields(self) -> None:
        unlabelled: list[dict] = []
        field_idx = 0

        for ff_data in self.doc.element.body.iter(qn("w:ffData")):
            name = ff_data.get(qn("w:name"), "") or ""
            if not str(name).strip():
                unlabelled.append({"excerpt": f"Form field {field_idx + 1}"})
            field_idx += 1

        if not unlabelled:
            return

        self.add_violation(
            violation_id="docx_form_no_name_001",
            checkpoint_id="GIGW_5.2.45",
            standard=Standard.WCAG_2_1_AA,
            severity=Severity.SERIOUS,
            category="form_fields",
            wcag_criterion="3.3.2",
            description=(
                f"{len(unlabelled)} of {field_idx} form field(s) have no name, so "
                "assistive technology has nothing to announce when focus reaches them."
            ),
            location=f"{len(unlabelled)} form field(s) in the document body",
            impact=(
                "A screen reader user tabbing through this form hears \"edit box\" with "
                "no indication of what to type. Any visible label printed beside the "
                "field is not connected to it in the file, so it is never announced."
            ),
            remediation=(
                "Give every field a name that matches its visible label, using content "
                "controls rather than legacy form fields where possible."
            ),
            fix_steps=[
                (
                    "Turn on the Developer tab: File → Options → Customise Ribbon → tick "
                    "Developer."
                ),
                (
                    "Click each field, then Developer → Properties, and set the Title to "
                    "the same wording as the visible label."
                ),
                (
                    "Prefer modern content controls over legacy form fields — legacy "
                    "fields expose far less to assistive technology."
                ),
                (
                    "Where a field needs a format, say so in the label itself, for "
                    'example "Date of birth (DD/MM/YYYY)" or '
                    '"Mobile number (+91 98765 43210)".'
                ),
                (
                    "Tab through the finished form with a screen reader to confirm every "
                    "field announces its label."
                ),
            ],
            occurrences=len(unlabelled),
            occurrence_list=unlabelled[:MAX_LISTED_OCCURRENCES],
        )
