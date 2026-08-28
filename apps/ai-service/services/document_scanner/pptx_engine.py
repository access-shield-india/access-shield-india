"""PPTX accessibility scanning engine for the Document Scanner."""

import logging
from typing import Optional

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from services.document_scanner.base import (
    BaseDocumentEngine,
    Severity,
    Standard,
)

logger = logging.getLogger(__name__)

GENERIC_LINK_TEXT = frozenset(
    {"click here", "here", "read more", "more", "link", "this link", "click", "www"}
)
PML_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
XDR_NS = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
PICTURE_SHAPE_TYPES = {MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE}

#: Minimum body text size for projected slides, per GIGW 5.1.15 guidance.
MIN_BODY_FONT_PT = 18

#: Relative luminance above which text is a likely contrast failure on white.
LIGHT_TEXT_LUMINANCE = 0.45


def _relative_luminance(rgb) -> float:
    """Rough relative luminance from an RGBColor, for contrast triage."""
    try:
        red, green, blue = rgb[0] / 255.0, rgb[1] / 255.0, rgb[2] / 255.0
    except (TypeError, IndexError):
        return 0.0
    return 0.2126 * red + 0.7152 * green + 0.0722 * blue


class PptxAccessibilityEngine(BaseDocumentEngine):
    """Runs WCAG 2.1 AA and GIGW 3.0 checks against a PowerPoint presentation."""

    async def run_all_checks(self) -> list:
        checks = [
            ("presentation_title", self.check_presentation_title),
            ("slide_titles", self.check_slide_titles),
            ("reading_order", self.check_reading_order),
            ("image_alt_text", self.check_image_alt_text),
            ("colour_contrast", self.check_colour_contrast),
            ("animation_timing", self.check_animation_timing),
            ("font_size", self.check_font_size),
            ("hyperlinks", self.check_hyperlinks),
        ]
        try:
            self.prs = Presentation(self.file_path)
        except Exception as e:
            logger.error("Failed to open PPTX: %s", e, exc_info=True)
            self.add_scan_error("document_load", e)
            return self.violations

        for name, check_fn in checks:
            try:
                await check_fn()
            except Exception as e:
                logger.error("PPTX check '%s' failed: %s", name, e, exc_info=True)
                self.add_scan_error(name, e)
        return self.violations

    # ── Helpers ─────────────────────────────────────────────────────────────

    def _slide_count(self) -> int:
        try:
            return len(self.prs.slides)
        except Exception:
            return 0

    @staticmethod
    def _slide_title(slide) -> str:
        """Slide title text, used to name a location the author will recognise."""
        try:
            title_shape = slide.shapes.title
            if title_shape is not None:
                return (title_shape.text or "").strip()
        except Exception:
            pass
        return ""

    def _slide_label(self, slide_num: int, slide) -> str:
        title = self._slide_title(slide)
        return f'Slide {slide_num} ("{title[:60]}")' if title else f"Slide {slide_num}"

    # ── Checks ──────────────────────────────────────────────────────────────

    async def check_presentation_title(self) -> None:
        title = self.prs.core_properties.title
        if title and str(title).strip():
            return

        first_title = ""
        try:
            for slide in self.prs.slides:
                first_title = self._slide_title(slide)
                if first_title:
                    break
        except Exception:
            first_title = ""

        self.add_violation(
            violation_id="pptx_no_pres_title_001",
            checkpoint_id="GIGW_5.2.28",
            standard=Standard.WCAG_2_1_AA,
            severity=Severity.SERIOUS,
            category="metadata",
            wcag_criterion="2.4.2",
            description=(
                "The presentation has no Title in its file properties."
                + (
                    f' Its first slide is titled "{first_title[:80]}".'
                    if first_title
                    else ""
                )
            ),
            location="File → Info → Properties → Title",
            excerpt=first_title,
            impact=(
                "Screen readers announce the file name or 'Untitled Presentation' when "
                "the deck opens, so a reader cannot confirm they have the right file. "
                "The empty Title also carries into any PDF exported from this deck."
            ),
            remediation=(
                "Set the Title in file properties. In PowerPoint: File → Info → "
                "Properties → Title."
            ),
            fix_steps=[
                "In PowerPoint, open File → Info.",
                "In the Properties panel, click Title.",
                (
                    "Enter the presentation's full name"
                    + (
                        f' — "{first_title[:80]}" looks right.'
                        if first_title
                        else ", as it would appear on an agenda."
                    )
                ),
                "Save the file.",
            ],
            auto_fixable=True,
        )

    async def check_slide_titles(self) -> None:
        seen_titles: dict[str, int] = {}

        for slide_num, slide in enumerate(self.prs.slides, 1):
            title_shape = slide.shapes.title
            body_hint = self._first_body_text(slide)

            if title_shape is None:
                self.add_violation(
                    violation_id=f"pptx_no_title_placeholder_s{slide_num}",
                    checkpoint_id="GIGW_5.2.28",
                    standard=Standard.WCAG_2_1_AA,
                    severity=Severity.SERIOUS,
                    category="slide_titles",
                    wcag_criterion="2.4.2",
                    description=(
                        f"Slide {slide_num} has no title placeholder, so it has no title "
                        "for assistive technology to announce."
                        + (f' Its content begins "{body_hint[:60]}".' if body_hint else "")
                    ),
                    location=f"Slide {slide_num}",
                    slide=slide_num,
                    excerpt=body_hint,
                    impact=(
                        "Screen reader users move through a deck by slide title, and the "
                        "Outline view and any generated table of contents both rely on "
                        "it. A slide with no title placeholder is unreachable that way — "
                        "a title typed into a plain text box does not count."
                    ),
                    remediation=(
                        "Apply a slide layout that includes a title placeholder, then add "
                        "the title text. If the title should not be visible, keep the "
                        "placeholder and move it off-slide rather than deleting it."
                    ),
                    fix_steps=[
                        f"Go to slide {slide_num}.",
                        (
                            "Choose Home → Layout and pick a layout that includes a Title "
                            "placeholder."
                        ),
                        (
                            "Type a title that describes what this slide covers"
                            + (
                                f', for example something based on "{body_hint[:50]}".'
                                if body_hint
                                else "."
                            )
                        ),
                        (
                            "For a deliberately title-free slide, such as a full-bleed "
                            "image, keep the placeholder, add the title, then drag it "
                            "outside the slide boundary. It stays available to screen "
                            "readers without appearing on screen."
                        ),
                        (
                            "Confirm every slide now has a title in View → Outline View."
                        ),
                    ],
                )
                continue

            title_text = (title_shape.text or "").strip()

            if not title_text:
                self.add_violation(
                    violation_id=f"pptx_empty_title_s{slide_num}",
                    checkpoint_id="GIGW_5.2.28",
                    standard=Standard.WCAG_2_1_AA,
                    severity=Severity.SERIOUS,
                    category="slide_titles",
                    wcag_criterion="2.4.2",
                    description=(
                        f"Slide {slide_num} has a title placeholder but it is empty."
                        + (f' Its content begins "{body_hint[:60]}".' if body_hint else "")
                    ),
                    location=f"Slide {slide_num}",
                    slide=slide_num,
                    excerpt=body_hint,
                    impact=(
                        "The slide appears in a screen reader's slide list as a blank "
                        "entry, and in Outline view as an empty line. Readers navigating "
                        "by title cannot tell what is on it or whether it matters."
                    ),
                    remediation=(
                        f"Add title text to slide {slide_num} describing what the slide "
                        "covers."
                    ),
                    fix_steps=[
                        f"Go to slide {slide_num} and click into the title placeholder.",
                        (
                            "Type a title that says what this slide shows, not a generic "
                            "label such as 'Overview'."
                        ),
                        (
                            "Keep it unique across the deck so readers can tell slides "
                            "apart from the title alone."
                        ),
                    ],
                    auto_fixable=False,
                )
                continue

            if title_text in seen_titles:
                first_seen = seen_titles[title_text]
                self.add_violation(
                    violation_id=f"pptx_duplicate_title_s{slide_num}",
                    checkpoint_id="GIGW_5.2.28",
                    standard=Standard.WCAG_2_1_AA,
                    severity=Severity.MODERATE,
                    category="slide_titles",
                    wcag_criterion="2.4.2",
                    description=(
                        f'Slide {slide_num} repeats the title "{title_text[:60]}", '
                        f"already used on slide {first_seen}."
                    ),
                    location=f"Slide {slide_num}",
                    slide=slide_num,
                    excerpt=title_text,
                    impact=(
                        "A screen reader user choosing from a list of slide titles sees "
                        "several identical entries and cannot tell which is which, so "
                        "they have to open each in turn to find the content they want."
                    ),
                    remediation=(
                        f'Make the title on slide {slide_num} distinct from slide '
                        f"{first_seen}, for example by adding the part or stage it covers."
                    ),
                    fix_steps=[
                        f"Go to slide {slide_num}.",
                        (
                            f'Change the title from "{title_text[:60]}" to something that '
                            "distinguishes it, such as adding a stage, region or period."
                        ),
                        (
                            "For a topic that genuinely spans several slides, number them: "
                            f'"{title_text[:40]} (1 of 3)".'
                        ),
                    ],
                )
            else:
                seen_titles[title_text] = slide_num

    @staticmethod
    def _first_body_text(slide) -> str:
        """First non-title text on a slide, to describe what it contains."""
        try:
            title_shape = slide.shapes.title
            for shape in slide.shapes:
                if shape is title_shape or not shape.has_text_frame:
                    continue
                text = (shape.text_frame.text or "").strip()
                if text:
                    return " ".join(text.split())
        except Exception:
            pass
        return ""

    async def check_reading_order(self) -> None:
        for slide_num, slide in enumerate(self.prs.slides, 1):
            shapes = [s for s in slide.shapes if s.has_text_frame or self._is_picture_shape(s)]
            if len(shapes) < 3:
                continue

            # Compare the order shapes are stored in (which drives narration)
            # against top-to-bottom visual order.
            visual_order = sorted(range(len(shapes)), key=lambda i: (shapes[i].top or 0))
            if visual_order == list(range(len(shapes))):
                continue

            first_out_of_place = next(
                (
                    position
                    for position, index in enumerate(visual_order)
                    if position != index
                ),
                0,
            )
            shape_name = getattr(shapes[visual_order[first_out_of_place]], "name", "")

            self.add_violation(
                violation_id=f"pptx_reading_order_s{slide_num}",
                checkpoint_id="GIGW_5.2.8",
                standard=Standard.WCAG_2_1_AA,
                severity=Severity.SERIOUS,
                category="reading_order",
                wcag_criterion="1.3.2",
                location=self._slide_label(slide_num, slide),
                slide=slide_num,
                excerpt=shape_name,
                description=(
                    f"On slide {slide_num} the stored order of the "
                    f"{len(shapes)} content shapes does not match their top-to-bottom "
                    "layout, so narration order may differ from what a viewer sees."
                ),
                impact=(
                    "Screen readers narrate shapes in the order they are stored, not the "
                    "order they appear. A reader can hear a caption before the chart it "
                    "describes, or a conclusion before the data behind it, with no "
                    "indication that the sequence is wrong."
                ),
                remediation=(
                    "Reorder the shapes in the Selection Pane so narration follows the "
                    "intended reading sequence, starting with the title."
                ),
                fix_steps=[
                    f"Go to slide {slide_num} and open Home → Arrange → Selection Pane.",
                    (
                        "Note that the pane lists shapes back to front: the item at the "
                        "bottom of the list is narrated first."
                    ),
                    (
                        "Drag the title to the bottom of the list, then arrange the rest "
                        "so reading upward matches the order you want them heard."
                    ),
                    (
                        "Check the result with the Accessibility Checker "
                        "(Review → Check Accessibility → Reading Order)."
                    ),
                    (
                        "Remove empty placeholders and stray text boxes rather than "
                        "leaving them in the order."
                    ),
                ],
            )

    def _get_shape_alt_text(self, shape) -> str:
        descr = shape.element.get(f"{{{XDR_NS}}}descr", "") or ""

        if not descr:
            for elem in shape.element.iter():
                if elem.tag.endswith("}cNvPr"):
                    descr = elem.get("descr", "") or ""
                    break
        return descr

    def _is_picture_shape(self, shape) -> bool:
        try:
            if shape.shape_type in PICTURE_SHAPE_TYPES:
                return True
        except Exception:
            pass
        for elem in shape.element.iter():
            if elem.tag.endswith("}pic"):
                return True
        return False

    async def check_image_alt_text(self) -> None:
        for slide_num, slide in enumerate(self.prs.slides, 1):
            for shape in slide.shapes:
                if not self._is_picture_shape(shape):
                    continue

                descr = self._get_shape_alt_text(shape)
                if descr.strip():
                    continue

                slide_title = self._slide_title(slide)
                self.add_violation(
                    violation_id=f"pptx_no_alt_s{slide_num}_sh{shape.shape_id}",
                    checkpoint_id="GIGW_5.2.1",
                    standard=Standard.WCAG_2_1_AA,
                    severity=Severity.CRITICAL,
                    category="alt_text",
                    wcag_criterion="1.1.1",
                    location=f"{self._slide_label(slide_num, slide)}, shape \"{shape.name}\"",
                    slide=slide_num,
                    excerpt=shape.name,
                    heading_path=[slide_title] if slide_title else None,
                    description=(
                        f'The image "{shape.name}" on slide {slide_num} has no '
                        "alternative text."
                        + (f' The slide is titled "{slide_title[:60]}".' if slide_title else "")
                    ),
                    impact=(
                        "Blind readers get nothing from this image. On a slide where the "
                        "image carries the argument — a chart, a process diagram, a "
                        "screenshot — the point of the slide is lost entirely."
                    ),
                    remediation=(
                        "Right-click the image → Edit Alt Text and describe what it tells "
                        "the audience, or mark it decorative if it carries no information."
                    ),
                    fix_steps=[
                        f'Go to slide {slide_num} and select the image "{shape.name}".',
                        "Right-click it and choose Edit Alt Text.",
                        (
                            "Write what the audience is meant to take from the image. For "
                            "a chart, state the finding rather than describing the bars."
                        ),
                        (
                            "If the image is a background texture, divider or decorative "
                            "logo, tick 'Mark as decorative' instead."
                        ),
                        (
                            "For a complex diagram, put the full explanation in the "
                            "speaker notes and reference it from the alt text — alt text "
                            "should stay under about two sentences."
                        ),
                    ],
                    auto_fixable=False,
                )

    async def check_colour_contrast(self) -> None:
        for slide_num, slide in enumerate(self.prs.slides, 1):
            light_runs: list[str] = []

            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        rgb = self._run_rgb(run)
                        if rgb is None:
                            continue
                        if _relative_luminance(rgb) > LIGHT_TEXT_LUMINANCE:
                            text = (run.text or "").strip()
                            if text:
                                light_runs.append(text)

            if not light_runs:
                continue

            self.add_violation(
                violation_id=f"pptx_low_contrast_s{slide_num}",
                checkpoint_id="GIGW_5.2.14",
                standard=Standard.WCAG_2_1_AA,
                severity=Severity.SERIOUS,
                category="colour_contrast",
                wcag_criterion="1.4.3",
                description=(
                    f"Slide {slide_num} has {len(light_runs)} run(s) of light-coloured "
                    "text that are likely to fall below the 4.5:1 contrast minimum "
                    f'against a light background. For example: "{light_runs[0][:60]}".'
                ),
                location=self._slide_label(slide_num, slide),
                slide=slide_num,
                excerpt=light_runs[0],
                impact=(
                    "Low-contrast text on a projected slide is the most common reason "
                    "audiences cannot read a deck. Ambient light, older projectors and "
                    "low vision all compound the problem, and nobody in the room can fix "
                    "it while the slide is on screen."
                ),
                remediation=(
                    "Measure the contrast of this text against its real background and "
                    "darken it until body text reaches 4.5:1, or 3:1 for text at 18pt "
                    "and above."
                ),
                fix_steps=[
                    f"Go to slide {slide_num} and select the light-coloured text.",
                    (
                        "Measure it against its actual background with the free TPGi "
                        "Colour Contrast Analyser."
                    ),
                    (
                        "Darken the text, or add a solid panel behind it if it sits over "
                        "an image or gradient."
                    ),
                    (
                        "Fix it in the Slide Master rather than slide by slide if the "
                        "colour comes from the theme (View → Slide Master)."
                    ),
                    (
                        "Aim above the minimum for projected content — 7:1 survives a "
                        "bright room far better than 4.5:1."
                    ),
                ],
            )

    @staticmethod
    def _run_rgb(run):
        try:
            colour = run.font.color
            if colour is None or colour.rgb is None:
                return None
            return colour.rgb
        except Exception:
            return None

    async def check_animation_timing(self) -> None:
        for slide_num, slide in enumerate(self.prs.slides, 1):
            transition = slide.element.find(f".//{{{PML_NS}}}transition")
            if transition is None:
                continue

            adv_tm = transition.get("advTm")
            if adv_tm is None:
                continue
            try:
                adv_ms = int(adv_tm)
            except ValueError:
                continue
            if adv_ms <= 0:
                continue

            seconds = adv_ms // 1000
            self.add_violation(
                violation_id=f"pptx_auto_advance_s{slide_num}",
                checkpoint_id="GIGW_5.2.25",
                standard=Standard.WCAG_2_1_AA,
                severity=Severity.SERIOUS,
                category="animation_timing",
                wcag_criterion="2.2.2",
                description=(
                    f"Slide {slide_num} advances automatically after {seconds} second(s), "
                    "so the reader cannot control how long it stays on screen."
                ),
                location=self._slide_label(slide_num, slide),
                slide=slide_num,
                excerpt=f"advances after {seconds}s",
                impact=(
                    "Readers who need longer — anyone using a screen reader or magnifier, "
                    "anyone reading in a second language, anyone with a cognitive "
                    "disability — lose the slide mid-sentence with no way to get it back. "
                    "WCAG 2.2.2 requires timed content to be pausable when it runs beyond "
                    "five seconds."
                ),
                remediation=(
                    "Remove the automatic advance so slides move on click, or provide a "
                    "pause control if the deck must run unattended."
                ),
                fix_steps=[
                    f"Go to slide {slide_num} and open the Transitions tab.",
                    (
                        "In the Timing group, untick 'After' and tick 'On Mouse Click'."
                    ),
                    (
                        "Apply the change to the whole deck with 'Apply To All' once it "
                        "is right."
                    ),
                    (
                        "If the deck genuinely has to run unattended, such as on a kiosk, "
                        "keep the timing but publish a self-paced copy alongside it and "
                        "make sure the kiosk exposes pause and back controls."
                    ),
                ],
                auto_fixable=True,
            )

    async def check_font_size(self) -> None:
        for slide_num, slide in enumerate(self.prs.slides, 1):
            total_runs = 0
            small_runs = 0
            smallest: Optional[float] = None
            sample = ""

            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        if run.font.size is None:
                            continue
                        size_pt = run.font.size.pt
                        total_runs += 1
                        if size_pt < MIN_BODY_FONT_PT:
                            small_runs += 1
                            if smallest is None or size_pt < smallest:
                                smallest = size_pt
                                sample = (run.text or "").strip()

            if total_runs == 0 or (small_runs / total_runs) <= 0.10:
                continue

            self.add_violation(
                violation_id=f"pptx_small_font_s{slide_num}",
                checkpoint_id="GIGW_5.1.15",
                standard=Standard.GIGW_3_0,
                severity=Severity.MINOR,
                category="readability",
                wcag_criterion="1.4.4",
                description=(
                    f"Slide {slide_num} has {small_runs} of {total_runs} text run(s) "
                    f"below {MIN_BODY_FONT_PT}pt"
                    + (f", the smallest at {smallest:g}pt" if smallest else "")
                    + (f' (\"{sample[:50]}\")' if sample else "")
                    + "."
                ),
                location=self._slide_label(slide_num, slide),
                slide=slide_num,
                excerpt=sample,
                impact=(
                    "Text below 18pt is unreadable from the back of a room and hard to "
                    "read on a laptop when the deck is shared as a file. Small text is "
                    "usually a sign the slide is carrying more content than a slide can."
                ),
                remediation=(
                    f"Raise body text to at least {MIN_BODY_FONT_PT}pt and titles to "
                    "24pt or more, splitting the slide if the content no longer fits."
                ),
                fix_steps=[
                    f"Go to slide {slide_num} and select the small text.",
                    (
                        f"Raise body text to at least {MIN_BODY_FONT_PT}pt and titles to "
                        "24pt or above."
                    ),
                    (
                        "If the content no longer fits, split it across two slides rather "
                        "than shrinking it back down."
                    ),
                    (
                        "Move detail, sources and caveats into the speaker notes, which "
                        "screen readers can access and which do not need to be legible "
                        "from a distance."
                    ),
                    (
                        "Set the sizes in View → Slide Master so later edits inherit them."
                    ),
                ],
            )

    async def check_hyperlinks(self) -> None:
        link_idx = 0
        for slide_num, slide in enumerate(self.prs.slides, 1):
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        if not run.hyperlink or not run.hyperlink.address:
                            continue

                        link_text = (run.text or "").strip()
                        address = run.hyperlink.address or ""
                        lowered = link_text.lower()
                        is_generic = lowered in GENERIC_LINK_TEXT
                        is_bare_url = lowered.startswith(("http://", "https://", "www."))

                        if is_generic or is_bare_url:
                            reason = (
                                "does not say where it goes"
                                if is_generic
                                else "shows the raw web address"
                            )
                            self.add_violation(
                                violation_id=f"pptx_generic_link_s{slide_num}_{link_idx}",
                                checkpoint_id="GIGW_5.2.30",
                                standard=Standard.WCAG_2_1_AA,
                                severity=(
                                    Severity.SERIOUS if is_generic else Severity.MODERATE
                                ),
                                category="link_text",
                                wcag_criterion="2.4.4",
                                description=(
                                    f'A link on slide {slide_num} reads "{link_text[:60]}", '
                                    f"which {reason}."
                                    + (f" It points to {address[:80]}." if address else "")
                                ),
                                location=self._slide_label(slide_num, slide),
                                slide=slide_num,
                                excerpt=link_text,
                                impact=(
                                    "Screen reader users often review the links on a slide "
                                    "as a list, out of context. A raw address is read out "
                                    "character by character, and a list of identical "
                                    '"click here" entries gives no way to choose.'
                                ),
                                remediation=(
                                    "Replace the link text with the name of the page or "
                                    "document it opens, keeping the same target."
                                ),
                                fix_steps=[
                                    f"Go to slide {slide_num} and select the link.",
                                    (
                                        "Right-click → Edit Link (or press Ctrl+K) and "
                                        "leave the Address unchanged."
                                    ),
                                    (
                                        "Replace 'Text to display' with the destination's "
                                        'name, for example "GIGW 3.0 guidelines (PDF)".'
                                    ),
                                    (
                                        "If the audience needs the address for a printed "
                                        "handout, put it in the speaker notes or on a "
                                        "references slide."
                                    ),
                                ],
                            )
                        link_idx += 1
