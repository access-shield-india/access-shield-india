"""
Standalone smoke check for the DOCX engine's location and grouping behaviour.

Run from apps/ai-service:
    PYTHONPATH=. python tests/smoke_docx_report.py

Kept out of the pytest suite because it builds fixtures on the fly and prints a
human-readable dump rather than asserting on exact strings.
"""

import asyncio
import importlib.util
import json
import sys
import types
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import RGBColor

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


def _load_engine_modules():
    """
    Load the scanner modules directly.

    `services/__init__.py` eagerly imports the whole AI service graph (Redis,
    Anthropic, settings), none of which the document engines need. Stubbing the
    parent packages keeps this check dependency-light.
    """
    for package in ("services", "services.document_scanner"):
        stub = types.ModuleType(package)
        stub.__path__ = []  # type: ignore[attr-defined]
        sys.modules[package] = stub

    def load(name: str, relative_path: str):
        spec = importlib.util.spec_from_file_location(name, ROOT / relative_path)
        module = importlib.util.module_from_spec(spec)
        sys.modules[name] = module
        spec.loader.exec_module(module)
        return module

    load("services.document_scanner.base", "services/document_scanner/base.py")
    load("services.document_scanner.locator", "services/document_scanner/locator.py")
    return load(
        "services.document_scanner.docx_engine", "services/document_scanner/docx_engine.py"
    )


DocxAccessibilityEngine = _load_engine_modules().DocxAccessibilityEngine


def build_fixture(path: str) -> None:
    """A document with the defects the report needs to describe precisely."""
    doc = Document()

    doc.add_heading("Accessibility Services Proposal", level=1)
    doc.add_paragraph(
        "This proposal sets out scope, timelines and commercials for the "
        "engagement. " * 12
    )

    doc.add_heading("Scope of Work", level=1)
    doc.add_paragraph("The scope covers four workstreams. " * 20)

    # Heading level skip: H1 -> H3.
    doc.add_heading("Workstream Detail", level=3)
    doc.add_paragraph("Detail of each workstream follows. " * 18)

    # Empty heading.
    doc.add_heading("", level=2)
    doc.add_paragraph("Timelines are indicative and subject to sign-off. " * 15)

    # Selective colour emphasis — should be flagged and grouped by colour.
    for index in range(6):
        para = doc.add_paragraph("Deliverable status for milestone ")
        run = para.add_run(f"M{index + 1} is at risk")
        run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        para.add_run(" and will be reviewed at the next checkpoint.")

    # Whole-paragraph colour — styling, should NOT be flagged.
    styled = doc.add_paragraph()
    styled_run = styled.add_run("Confidential — for internal circulation only.")
    styled_run.font.color.rgb = RGBColor(0x6D, 0x28, 0xD9)

    # Highlighting — always flagged.
    for index in range(3):
        para = doc.add_paragraph("Clause ")
        run = para.add_run(f"{index + 4}.2 has been amended")
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        para.add_run(" since the previous revision.")

    # Manual bullets.
    for item in ("Discovery workshop", "Audit and reporting", "Remediation support"):
        doc.add_paragraph(f"- {item}")

    # Table without a header row.
    table = doc.add_table(rows=3, cols=3)
    headers = ("Milestone", "Date", "Fee (INR)")
    for col, value in enumerate(headers):
        table.rows[0].cells[col].text = value
    for row in range(1, 3):
        for col in range(3):
            table.rows[row].cells[col].text = f"r{row}c{col}"

    doc.add_paragraph("Annexures follow. " * 30)

    doc.save(path)


async def main() -> None:
    fixture = "/tmp/smoke_proposal.docx"
    build_fixture(fixture)

    engine = DocxAccessibilityEngine(fixture)
    violations = await engine.run_all_checks()

    print(f"\nDistinct findings: {len(violations)}")
    print(
        "Locator pagination: "
        f"{'real' if engine.locator and engine.locator.uses_real_pagination else 'estimated'}"
    )
    print("=" * 78)

    for violation in violations:
        data = violation.to_dict()
        print(f"\n[{data['severity'].upper()}] {data['description']}")
        print(f"  standard      : {data['standard']}  checkpoint={data['checkpoint_id']}")
        print(f"  wcag          : {data['wcag_criterion']}")
        print(f"  location      : {data['location']}")
        print(
            "  anchors       : "
            f"page={data['page']} line={data['line']} para={data['paragraph']}"
        )
        print(f"  heading_path  : {data['heading_path']}")
        print(f"  excerpt       : {data['excerpt']}")
        print(f"  occurrences   : {data['occurrences']}")
        if data["occurrence_list"]:
            for entry in data["occurrence_list"][:4]:
                print(f"      - {json.dumps(entry, ensure_ascii=False)}")
        print("  fix_steps     :")
        for step_number, step in enumerate(data["fix_steps"] or [], 1):
            print(f"      {step_number}. {step}")

    # Sanity assertions on the behaviour the report depends on.
    colour_findings = [v for v in violations if v.category == "colour_contrast"]
    assert colour_findings, "expected selective colour emphasis to be flagged"
    assert all(
        v.occurrences == len(v.occurrence_list) or v.occurrences >= 1
        for v in colour_findings
    )
    assert len(colour_findings) <= 3, (
        f"colour findings should be grouped per colour, got {len(colour_findings)}"
    )

    assert any(v.category == "table_structure" for v in violations), "table check missing"
    assert any(
        v.violation_id.startswith("docx_heading_skip") for v in violations
    ), "heading skip not detected"
    assert any(
        v.violation_id == "docx_empty_heading_001" for v in violations
    ), "empty heading not detected"
    assert any(
        v.violation_id == "docx_manual_list_001" for v in violations
    ), "manual list not detected"

    titled = [v for v in violations if v.violation_id == "docx_no_title_001"]
    assert titled, "missing-title finding expected"
    assert "Accessibility Services Proposal" in titled[0].description, (
        "title finding should quote the document's own headline"
    )

    for violation in violations:
        assert violation.fix_steps, f"{violation.violation_id} has no fix steps"

    print("\n" + "=" * 78)
    print("All smoke assertions passed.")


if __name__ == "__main__":
    asyncio.run(main())
