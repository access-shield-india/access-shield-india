"""Generate test fixture files for document scanner engine tests."""

import io
from pathlib import Path

from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import PatternFill
from PIL import Image
from pptx import Presentation
from pptx.util import Inches as PptxInches, Pt

FIXTURES_DIR = Path(__file__).parent


def _make_png_bytes() -> io.BytesIO:
    img = Image.new("RGB", (50, 50), color="red")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return buf


def generate_no_headings_docx() -> None:
    doc = Document()
    doc.add_paragraph(" ".join(["accessibility"] * 600))
    doc.save(FIXTURES_DIR / "test_no_headings.docx")


def generate_heading_skip_docx() -> None:
    doc = Document()
    doc.add_paragraph("Main Section", style="Heading 1")
    doc.add_paragraph("Skipped level subsection", style="Heading 3")
    doc.save(FIXTURES_DIR / "test_heading_skip.docx")


def generate_images_no_alt_docx() -> None:
    doc = Document()
    doc.add_paragraph("Document with image")
    doc.add_picture(_make_png_bytes(), width=Inches(1))
    doc.save(FIXTURES_DIR / "test_images_no_alt.docx")


def generate_good_structure_docx() -> None:
    doc = Document()
    doc.core_properties.title = "Accessible Test Document"
    doc.core_properties.language = "en-IN"
    doc.add_paragraph("Introduction", style="Heading 1")
    doc.add_paragraph("This is accessible body content with proper structure.")
    doc.add_paragraph("Details", style="Heading 2")
    doc.add_paragraph("More content here.")
    doc.save(FIXTURES_DIR / "test_good_structure.docx")


def generate_missing_titles_pptx() -> None:
    prs = Presentation()
    blank_layout = prs.slide_layouts[6]
    prs.slides.add_slide(blank_layout)
    prs.slides.add_slide(blank_layout)
    prs.save(FIXTURES_DIR / "test_missing_titles.pptx")


def generate_auto_advance_pptx() -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Auto Advance Slide"
    pml_ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    transition = slide.element.find(f".//{{{pml_ns}}}transition")
    if transition is None:
        from lxml import etree

        transition = etree.SubElement(slide.element, f"{{{pml_ns}}}transition")
    transition.set("advTm", "5000")
    prs.save(FIXTURES_DIR / "test_auto_advance.pptx")


def generate_colour_only_xlsx() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "StatusData"
    fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
    for row in range(2, 25):
        for col in range(2, 8):
            cell = ws.cell(row=row, column=col)
            cell.fill = fill
    wb.save(FIXTURES_DIR / "test_colour_only.xlsx")


def generate_chart_no_alt_xlsx() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "SalesData"
    for row, value in enumerate([10, 20, 30, 40], start=1):
        ws.cell(row=row, column=1, value=f"Q{row}")
        ws.cell(row=row, column=2, value=value)
    chart = BarChart()
    data = Reference(ws, min_col=2, min_row=1, max_row=4)
    cats = Reference(ws, min_col=1, min_row=2, max_row=4)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, "D2")
    wb.save(FIXTURES_DIR / "test_chart_no_alt.xlsx")


def generate_generic_sheet_names_xlsx() -> None:
    wb = Workbook()
    wb.active.title = "Sheet1"
    wb.create_sheet("Sheet2")
    wb.save(FIXTURES_DIR / "test_generic_sheet_names.xlsx")


def generate_all() -> None:
    FIXTURES_DIR.mkdir(parents=True, exist_ok=True)
    generate_no_headings_docx()
    generate_heading_skip_docx()
    generate_images_no_alt_docx()
    generate_good_structure_docx()
    generate_missing_titles_pptx()
    generate_auto_advance_pptx()
    generate_colour_only_xlsx()
    generate_chart_no_alt_xlsx()
    generate_generic_sheet_names_xlsx()
    print(f"Generated fixtures in {FIXTURES_DIR}")


if __name__ == "__main__":
    generate_all()
